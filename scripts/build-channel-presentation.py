"""Build a self-contained animated HTML deck, editable PPTX, PDF and preview.
Only existing project illustrations are used. No API inference or credentials.
"""
import argparse
import base64
import hashlib
import html
import json
import textwrap
import zipfile
from io import BytesIO
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.utils import ImageReader

W, H = 1600, 900
C = {'cream':'#F7F1E6','ink':'#182B31','dark':'#102E32','card':'#FFFFFF','muted':'#526368','amber':'#DEA84B','teal':'#1B6B65','light':'#ECF0E8','darkcard':'#1B4145'}


def font_path(bold=False):
    name = 'LiberationSans-Bold.ttf' if bold else 'LiberationSans-Regular.ttf'
    for folder in ['/usr/share/fonts/truetype/liberation2', '/usr/share/fonts/truetype/liberation']:
        p = Path(folder) / name
        if p.exists():
            return p
    raise RuntimeError('Install fonts-liberation before building the presentation')


def wrap(text, width, size, bold=False):
    font = ImageFont.truetype(str(font_path(bold)), size)
    lines = []
    for paragraph in text.split('\n'):
        line = ''
        for word in paragraph.split():
            if line and font.getlength(line + ' ' + word) > width:
                lines.append(line)
                line = word
            else:
                line += (' ' if line else '') + word
        lines.append(line)
    return lines


def scene(slide, number, total):
    dark = slide.get('dark', False)
    fg, muted, bg, card = ('#F7F1E6','#C3D5D0',C['dark'],C['darkcard']) if dark else (C['ink'],C['muted'],C['cream'],C['card'])
    elements = []

    def rect(x,y,w,h,color,r=18):
        elements.append({'kind':'rect','x':x,'y':y,'w':w,'h':h,'color':color,'r':r})

    def txt(x,y,w,value,size=30,color=None,bold=False,max_h=None):
        lines = wrap(value,w,size,bold)
        height = len(lines)*size*1.16
        if max_h and height > max_h:
            raise ValueError(f'Slide {number}: text exceeds its layout: {value[:50]} ({height}>{max_h})')
        if y+height > H-10:
            raise ValueError(f'Slide {number}: text outside canvas')
        elements.append({'kind':'text','x':x,'y':y,'w':w,'h':height+8,'lines':lines,'size':size,'color':color or fg,'bold':bold})

    def picture(key,x,y,w,h):
        elements.append({'kind':'image','asset':key,'x':x,'y':y,'w':w,'h':h})

    rect(0,0,W,H,bg,0)
    rect(72,63,48,7,C['amber'],3)
    txt(137,53,1370,slide['kicker'],20,muted,True)
    if slide['layout']=='cover':
        txt(72,145,960,slide['title'],94,bold=True,max_h=340)
        txt(75,523,870,slide['subtitle'],33,muted,max_h=120)
        rect(1040,160,488,520,C['cream'],28)
        picture(slide['image'],1070,183,428,466)
    else:
        txt(72,103,1456,slide['title'],66,bold=True,max_h=160)
        title_h=len(wrap(slide['title'],1456,66,True))*66*1.16
        txt(75,115+title_h,1430,slide['subtitle'],27,muted,max_h=70)
        top=304
        if slide['layout']=='cards':
            for i, item in enumerate(slide['cards']):
                x=72+i*493
                rect(x,top,469,365,card)
                txt(x+27,top+24,390,f'0{i+1}',25,C['amber'],True)
                txt(x+27,top+78,410,item['title'],35,bold=True,max_h=90)
                txt(x+27,top+179,410,item['body'],29,muted,max_h=170)
        elif slide['layout']=='metrics':
            for i, item in enumerate(slide['metrics']):
                x=72+i*371
                rect(x,top,343,365,card)
                txt(x+23,top+32,300,item['value'],48,C['teal'],True,max_h=120)
                txt(x+23,top+141,300,item['label'],30,bold=True,max_h=75)
                txt(x+23,top+229,300,item['body'],26,muted,max_h=115)
        elif slide['layout']=='flow':
            for i, item in enumerate(slide['cards']):
                x=72+i*371
                rect(x,top+8,343,345,card)
                rect(x+24,top+32,58,50,C['amber'],12)
                txt(x+43,top+38,50,str(i+1),29,C['ink'],True)
                txt(x+24,top+117,295,item['title'],34,bold=True,max_h=80)
                txt(x+24,top+205,295,item['body'],27,muted,max_h=130)
                if i<3:
                    txt(x+345,top+146,30,'›',43,C['amber'],True)
        elif slide['layout']=='split':
            for i,item in enumerate(slide['cards']):
                y=top+i*191
                rect(72,y,765,169,card)
                txt(99,y+23,700,item['title'],32,bold=True,max_h=42)
                txt(99,y+81,700,item['body'],28,muted,max_h=75)
            rect(877,top,651,360 if slide['image']!='eli' else 388,C['card'])
            picture(slide['image'],893,top+15,619,330 if slide['image']!='eli' else 356)
        elif slide['layout']=='compare':
            for i,item in enumerate(slide['images']):
                x=72+i*741
                rect(x,top-4,715,392,C['card'])
                picture(item['id'],x+16,top+9,683,243)
                txt(x+24,top+266,665,item['title'],27,C['ink'],True,max_h=66)
                txt(x+24,top+330,665,item['body'],23,C['muted'],max_h=58)
        elif slide['layout']=='table':
            for i,item in enumerate(slide['rows']):
                y=top+i*93
                rect(72,y,1456,79,card)
                txt(95,y+20,310,item['left'],29,C['teal'],True,max_h=36)
                txt(444,y+21,1050,item['right'],26,muted,max_h=42)
    rect(72,744,1456,75,C['amber'] if dark else '#E8E3D5',17)
    txt(96,766,1408,slide['takeaway'],26,C['ink'],True,max_h=35)
    txt(74,857,1250,'CANAL BÍBLICO EN LA NUBE  ·  Guía de réplica  ·  15 OCT 2026',17,muted)
    txt(1430,853,100,f'{number:02d} / {total:02d}',21,muted,True)
    return elements


def assets_from_zip(file, out):
    expected_zip='042df95b0a32b56223e10205f5cc0851515cbd4fba50ba53542d61f7d4b263b0'
    if hashlib.sha256(Path(file).read_bytes()).hexdigest()!=expected_zip:
        raise ValueError('Evidence archive checksum mismatch')
    cases=json.loads(Path('data/research/bible-angels-001/smolvlm-probe-cases.json').read_text())['cases']
    ids={'reader':'reader-control','wings':'wing-count','shirt':'shirt-collar'}
    result={'eli':Path('assets/references/style-v5-eli.png')}
    if hashlib.sha256(result['eli'].read_bytes()).hexdigest()!='9c3433c1a337fee080990a5be6613dc8adbc46296e2818477df82a7b63be2c07':
        raise ValueError('Approved project reference changed')
    folder=out/'source-assets';folder.mkdir(exist_ok=True)
    with zipfile.ZipFile(file) as archive:
        if archive.testzip() is not None:
            raise ValueError('Evidence ZIP corrupted')
        for key, case_id in ids.items():
            case=next(c for c in cases if c['id']==case_id)
            name=case_id+'-'+Path(case['path']).name
            data=archive.read(name)
            if hashlib.sha256(data).hexdigest()!=case['sha256']:
                raise ValueError('Source image checksum mismatch')
            path=folder/(key+'.jpg');path.write_bytes(data);result[key]=path
    return result


def contain(path, x,y,w,h):
    with Image.open(path) as im:
        scale=min(w/im.width,h/im.height)
        iw,ih=im.width*scale,im.height*scale
    return x+(w-iw)/2,y+(h-ih)/2,iw,ih


def build(data, output, assets):
    graphs=[scene(s,i+1,len(data['slides'])) for i,s in enumerate(data['slides'])]
    prs=Presentation();prs.core_properties.title=data['title'];prs.core_properties.author='youtube-autopilot';prs.core_properties.subject='Guía para crear un canal bíblico ilustrado en la nube';prs.slide_width=Inches(W/120);prs.slide_height=Inches(H/120)
    pdfmetrics.registerFont(TTFont('Deck',str(font_path())))
    pdfmetrics.registerFont(TTFont('DeckBold',str(font_path(True))))
    pdf=canvas.Canvas(str(output/'Guia.pdf'),pagesize=(W,H));pdf.setTitle(data['title']);pdf.setAuthor('youtube-autopilot')
    thumbs=[]
    for i,(slide,nodes) in enumerate(zip(data['slides'],graphs)):
        page=prs.slides.add_slide(prs.slide_layouts[6]);page.notes_slide.notes_text_frame.text=slide['notes']
        transition=OxmlElement('p:transition');transition.set('spd','med');transition.set('advClick','1');transition.append(OxmlElement('p:fade'));page._element.append(transition)
        im=Image.new('RGB',(W,H),C['cream']);draw=ImageDraw.Draw(im)
        for node in nodes:
            x,y,w,h=(node[k] for k in ['x','y','w','h'])
            if node['kind']=='rect':
                color=node['color'];shape=page.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if node['r'] else MSO_SHAPE.RECTANGLE,Inches(x/120),Inches(y/120),Inches(w/120),Inches(h/120));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color[1:]);shape.line.fill.background()
                if node['r']:shape.adjustments[0]=min(.2,node['r']/min(w,h))
                pdf.setFillColor(color);pdf.roundRect(x,H-y-h,w,h,node['r'],fill=1,stroke=0)
                draw.rounded_rectangle((x,y,x+w,y+h),radius=node['r'],fill=color)
            elif node['kind']=='image':
                path=assets[node['asset']];ix,iy,iw,ih=contain(path,x,y,w,h)
                page.shapes.add_picture(str(path),Inches(ix/120),Inches(iy/120),width=Inches(iw/120),height=Inches(ih/120))
                pdf.drawImage(ImageReader(str(path)),ix,H-iy-ih,width=iw,height=ih,mask='auto')
                source=Image.open(path).convert('RGB').resize((round(iw),round(ih)),Image.Resampling.LANCZOS);im.paste(source,(round(ix),round(iy)))
            else:
                size=node['size'];color=node['color'];font='DeckBold' if node['bold'] else 'Deck';tf=page.shapes.add_textbox(Inches(x/120),Inches(y/120),Inches(w/120),Inches(h/120)).text_frame;tf.clear();tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                for li,line in enumerate(node['lines']):
                    p=tf.paragraphs[0] if li==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size*.6);p.font.bold=node['bold'];p.font.color.rgb=RGBColor.from_string(color[1:]);p.space_before=p.space_after=Pt(0);p.line_spacing=1.16
                    pdf.setFont(font,size);pdf.setFillColor(color);pdf.drawString(x,H-y-size*.9-li*size*1.16,line)
                    draw.text((x,y+li*size*1.16),line,font=ImageFont.truetype(str(font_path(node['bold'])),size),fill=color,anchor='lt')
        thumb=im.resize((480,270),Image.Resampling.LANCZOS);thumbs.append(thumb);pdf.showPage()
    prs.save(output/'Presentacion.pptx');pdf.save()
    preview=Image.new('RGB',(4*480,4*300),'#E8E3D5');d=ImageDraw.Draw(preview)
    for i,thumb in enumerate(thumbs):
        x,y=i%4*480,i//4*300;preview.paste(thumb,(x,y));d.text((x+15,y+276),f'{i+1:02d}  {data["slides"][i]["kicker"]}',font=ImageFont.truetype(str(font_path()),15),fill=C['ink'])
    preview.save(output/'Vista-previa.jpg',quality=90)
    html_sections=[]
    for i,nodes in enumerate(graphs):
        content=[]
        for n,node in enumerate(nodes):
            style=f'left:{node["x"]/W*100}%;top:{node["y"]/H*100}%;width:{node["w"]/W*100}%;height:{node["h"]/H*100}%;--delay:{min(n*.035,.5)}s;'
            if node['kind']=='rect':
                content.append(f'<div class="node shape" style="{style}background:{node["color"]};border-radius:{node["r"]}px"></div>')
            elif node['kind']=='image':
                raw=Path(assets[node['asset']]).read_bytes();mime='image/png' if raw[:8]==b'\x89PNG\r\n\x1a\n' else 'image/jpeg';url='data:'+mime+';base64,'+base64.b64encode(raw).decode()
                content.append(f'<img class="node build illustration" alt="Ilustración original del proyecto" src="{url}" style="{style}">')
            else:
                size=node['size']/W*100
                content.append(f'<div class="node build text" style="{style}font-size:{size}cqw;color:{node["color"]};font-weight:{700 if node["bold"] else 400}">'+ '<br>'.join(html.escape(t) for t in node['lines'])+'</div>')
        html_sections.append(f'<section class="slide" role="group" aria-roledescription="diapositiva" aria-label="{i+1}. {html.escape(data["slides"][i]["title"])}">'+''.join(content)+'</section>')
    source_links=''.join(f'<li><a href="{html.escape(s["url"])}" target="_blank" rel="noopener noreferrer">{html.escape(s["label"])}</a></li>' for s in data['sources'])
    notes_json=json.dumps([s['notes'] for s in data['slides']],ensure_ascii=False).replace('</','<\\/')
    template='''<!doctype html><html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>De una idea a un canal en la nube</title><style>
:root{color-scheme:dark}*{box-sizing:border-box}body{margin:0;background:#0a171c;color:#edf3ee;font-family:Arial,sans-serif;display:flex;align-items:center;justify-content:center;min-height:100dvh;overflow:hidden}.shell{width:min(100vw,calc((100dvh - 78px)*16/9));container-type:inline-size;box-shadow:0 24px 100px #0008}.stage{aspect-ratio:16/9;position:relative;overflow:hidden;background:#f7f1e6}.slide{position:absolute;inset:0;opacity:0;visibility:hidden;transform:translateX(2%);transition:opacity .38s,transform .45s}.slide.active{opacity:1;visibility:visible;transform:none}.node{position:absolute}.text{line-height:1.16;white-space:nowrap}.illustration{object-fit:contain}.slide.active .build{animation:reveal .65s both;animation-delay:var(--delay)}@keyframes reveal{from{opacity:0;transform:translateY(12px)}to{opacity:1;transform:none}}nav{height:62px;display:flex;align-items:center;gap:9px;padding:8px 14px;background:#10282e;font-size:14px}button{border:1px solid #5d777b;background:#173a40;color:#fff;border-radius:9px;padding:9px 13px;cursor:pointer;font:inherit}button:hover,button:focus-visible{background:#27606a;outline:2px solid #dea84b}button:disabled{opacity:.4;cursor:default}#counter{margin-right:auto;font-variant-numeric:tabular-nums}.progress{height:4px;background:#29454c}.progress span{display:block;height:100%;background:#dea84b;transition:width .35s}dialog{max-width:850px;width:92vw;border:1px solid #789396;background:#10282e;color:#fff;border-radius:18px;padding:25px;line-height:1.55}dialog::backdrop{background:#000a}a{color:#f3c471}.help{opacity:.8}#notes p{font-size:18px}@media(max-width:700px){nav{height:auto;min-height:62px;flex-wrap:wrap;gap:6px;font-size:12px}button{padding:8px}.help{display:none}}@media(prefers-reduced-motion:reduce){.slide,.progress span{transition:none}.slide.active .build{animation:none}}@media print{body{display:block;background:white;overflow:visible}.shell{width:100%;box-shadow:none}.stage{overflow:visible;aspect-ratio:auto}.slide{position:relative;opacity:1;visibility:visible;transform:none;aspect-ratio:16/9;page-break-after:always}.build{animation:none!important}nav,.progress,dialog{display:none}}
</style></head><body><main class="shell"><div class="stage">__SLIDES__</div><div class="progress" aria-hidden="true"><span id="bar"></span></div><nav aria-label="Controles de presentación"><button id="prev" aria-label="Diapositiva anterior">←</button><span id="counter" aria-live="polite"></span><button id="next" aria-label="Diapositiva siguiente">→</button><button id="play">Reproducir</button><button id="full">Pantalla completa</button><button id="notesButton">Notas</button><button id="sourcesButton">Fuentes</button><span class="help">Flechas · Espacio · F</span></nav></main><dialog id="notes"><button class="close">Cerrar</button><h2>Notas del presentador</h2><p></p></dialog><dialog id="sources"><button class="close">Cerrar</button><h2>Fuentes y límites</h2><p>Guía verificada el 15 de octubre de 2026. Las conexiones propuestas no están activadas por abrir esta presentación. No contiene claves ni contraseñas.</p><ul>__SOURCES__</ul></dialog><script>
const slides=[...document.querySelectorAll('.slide')],notes=__NOTES__;let index=0,timer=null;const $=id=>document.getElementById(id);function show(n){index=Math.max(0,Math.min(slides.length-1,n));slides.forEach((s,i)=>{s.classList.toggle('active',i===index);s.setAttribute('aria-hidden',i!==index)});$('counter').textContent=`${index+1} / ${slides.length}`;$('bar').style.width=`${(index+1)/slides.length*100}%`;$('prev').disabled=index===0;$('next').disabled=index===slides.length-1;history.replaceState(null,'','#'+(index+1));$('notes').querySelector('p').textContent=notes[index];if(index===slides.length-1&&timer)togglePlay()};function togglePlay(){if(timer){clearInterval(timer);timer=null;$('play').textContent='Reproducir'}else{if(index===slides.length-1)show(0);timer=setInterval(()=>show(index+1),12000);$('play').textContent='Pausar'}}function full(){if(document.fullscreenElement)document.exitFullscreen();else document.documentElement.requestFullscreen?.()};$('prev').onclick=()=>show(index-1);$('next').onclick=()=>show(index+1);$('play').onclick=togglePlay;$('full').onclick=full;$('notesButton').onclick=()=>$('notes').showModal();$('sourcesButton').onclick=()=>$('sources').showModal();document.querySelectorAll('.close').forEach(b=>b.onclick=()=>b.closest('dialog').close());document.addEventListener('keydown',e=>{if(document.querySelector('dialog[open]'))return;if(e.key==='ArrowRight'){show(index+1);e.preventDefault()}if(e.key==='ArrowLeft'){show(index-1);e.preventDefault()}if(e.code==='Space'&&e.target===document.body){togglePlay();e.preventDefault()}if(e.key.toLowerCase()==='f')full();if(e.key==='Home')show(0);if(e.key==='End')show(slides.length-1)});let touchX=0;document.querySelector('.stage').addEventListener('touchstart',e=>touchX=e.changedTouches[0].screenX,{passive:true});document.querySelector('.stage').addEventListener('touchend',e=>{const dx=e.changedTouches[0].screenX-touchX;if(Math.abs(dx)>50)show(index+(dx<0?1:-1))},{passive:true});show((Number(location.hash.slice(1))||1)-1);
</script></body></html>'''
    (output/'ABRIR-PRESENTACION.html').write_text(template.replace('__SLIDES__',''.join(html_sections)).replace('__SOURCES__',source_links).replace('__NOTES__',notes_json))
    notes=['# Notas para presentar: '+data['title'],'','Fecha: '+data['date'],'']
    for i,s in enumerate(data['slides']):
        notes+=['## '+str(i+1)+'. '+s['title'].replace('\n',' '),'',s['notes'],'']
    notes+=['## Fuentes oficiales','']+[s['label']+': '+s['url'] for s in data['sources']]
    (output/'Notas-del-presentador.md').write_text('\n'.join(notes)+'\n')
    (output/'LEEME.txt').write_text('Abrí ABRIR-PRESENTACION.html con Chrome/Edge/Firefox: funciona sin servidor ni conexión. Flechas: navegar. Reproducir: avance cada 12 s. F: pantalla completa. Notas: explicación por diapositiva.\nPresentacion.pptx: editable; usar Presentación/F5 para ver transiciones de desvanecimiento.\nGuia.pdf: versión estática para enviar o imprimir.\nNo contiene credenciales. Ninguna cuenta se conecta ni publica al abrir los archivos.\n')
    names=['ABRIR-PRESENTACION.html','Presentacion.pptx','Guia.pdf','Vista-previa.jpg','Notas-del-presentador.md','LEEME.txt']
    manifest={'title':data['title'],'slides':len(data['slides']),'paidApiCalls':0,'liveConnectionsActivated':False,'files':[{'name':name,'bytes':(output/name).stat().st_size,'sha256':hashlib.sha256((output/name).read_bytes()).hexdigest()} for name in names]}
    (output/'manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    with zipfile.ZipFile(output/'Presentacion-Canal-Biblico.zip','w',zipfile.ZIP_DEFLATED) as archive:
        for name in names+['manifest.json']:
            archive.write(output/name,name)
    print(json.dumps({'event':'presentation_built','slides':len(data['slides']),'output':str(output),'files':manifest['files']}))


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--evidence-zip',required=True);parser.add_argument('--output',default='.cache/channel-presentation');args=parser.parse_args()
    output=Path(args.output);output.mkdir(parents=True,exist_ok=True)
    data=json.loads(Path('docs/presentations/channel-blueprint.json').read_text())
    assets=assets_from_zip(args.evidence_zip,output)
    build(data,output,assets)


if __name__=='__main__':
    main()
